import hashlib
import logging
import os
import tempfile
import uuid
import zipfile
from dataclasses import dataclass
from datetime import datetime, timedelta, timezone
from io import BytesIO
from pathlib import Path
from typing import BinaryIO

from docx.opc.exceptions import PackageNotFoundError as DocxPackageNotFoundError
from lxml.etree import XMLSyntaxError
from pypdf import PdfReader
from pypdf.errors import PdfReadError
from pptx.exc import PackageNotFoundError as PptxPackageNotFoundError
from sqlalchemy import select, update
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from app.core.config import settings
from app.core.errors import ApiError
from app.models.document import Document, DocumentChunk, ProcessingJob, Summary
from app.models.quiz import Quiz
from app.services.ai import AIProviderError, embed_document_chunks
from app.services.analytics import record_learning_event

logger = logging.getLogger(__name__)

MEDIA_TYPES = {
    ".pdf": "application/pdf",
    ".docx": (
        "application/vnd.openxmlformats-officedocument."
        "wordprocessingml.document"
    ),
    ".pptx": (
        "application/vnd.openxmlformats-officedocument."
        "presentationml.presentation"
    ),
}
OOXML_CONTENT = {
    ".docx": ("[Content_Types].xml", "word/document.xml"),
    ".pptx": ("[Content_Types].xml", "ppt/presentation.xml"),
}
CHUNK_MAX_CHARACTERS = 4000
INGEST_JOB = "DOCUMENT_INGEST"
DELETE_JOB = "DOCUMENT_DELETE"


class DocumentProcessingError(ValueError):
    pass


@dataclass(frozen=True)
class ExtractedChunk:
    chunk_index: int
    content: str
    page_number: int | None


@dataclass(frozen=True)
class ExtractedDocument:
    page_count: int | None
    chunks: list[ExtractedChunk]


def _safe_filename(filename: str) -> str:
    name = filename.replace("\\", "/").rsplit("/", maxsplit=1)[-1]
    name = name.replace("\x00", "").strip()
    if not name or name in {".", ".."}:
        raise ApiError(422, "VALIDATION_ERROR", "A valid filename is required.")
    return name[:255]


def _validate_zip(data: bytes, suffix: str) -> None:
    try:
        with zipfile.ZipFile(BytesIO(data)) as archive:
            members = archive.infolist()
            if not members or len(members) > 10_000:
                raise ValueError("Invalid document package.")
            if any(
                member.filename.startswith(("/", "\\"))
                or ".." in Path(member.filename.replace("\\", "/")).parts
                for member in members
            ):
                raise ValueError("Unsafe document package.")
            total_size = sum(member.file_size for member in members)
            if total_size > settings.document_max_uncompressed_bytes:
                raise ValueError("Document package exceeds extraction limits.")
            if any(
                member.file_size
                and member.compress_size == 0
                or (
                    member.compress_size > 0
                    and member.file_size / member.compress_size > 1000
                )
                for member in members
            ):
                raise ValueError("Document package exceeds extraction limits.")
            required = OOXML_CONTENT[suffix]
            if not set(required).issubset(archive.namelist()):
                raise ValueError("Unsupported OOXML package.")
    except (zipfile.BadZipFile, OSError, ValueError) as exc:
        raise ApiError(
            415,
            "UNSUPPORTED_FILE_TYPE",
            "The file type or contents are not supported.",
        ) from exc


def _validate_upload(filename: str, media_type: str, data: bytes) -> str:
    suffix = Path(filename.replace("\\", "/")).suffix.lower()
    expected_type = MEDIA_TYPES.get(suffix)
    if expected_type is None or media_type.lower() != expected_type:
        raise ApiError(
            415,
            "UNSUPPORTED_FILE_TYPE",
            "Upload a PDF, DOCX, or PPTX file with a matching content type.",
        )
    if not data:
        raise ApiError(
            415,
            "UNSUPPORTED_FILE_TYPE",
            "The uploaded file is empty or unreadable.",
        )
    if suffix == ".pdf":
        if not data.startswith(b"%PDF-"):
            raise ApiError(
                415,
                "UNSUPPORTED_FILE_TYPE",
                "The file type or contents are not supported.",
            )
    else:
        if not data.startswith((b"PK\x03\x04", b"PK\x05\x06", b"PK\x07\x08")):
            raise ApiError(
                415,
                "UNSUPPORTED_FILE_TYPE",
                "The file type or contents are not supported.",
            )
        _validate_zip(data, suffix)
    return expected_type


def _read_bounded(source: BinaryIO) -> bytes:
    content = source.read(settings.document_max_upload_bytes + 1)
    if len(content) > settings.document_max_upload_bytes:
        raise ApiError(
            413,
            "FILE_TOO_LARGE",
            "The uploaded file exceeds the configured size limit.",
        )
    return content


def _storage_path(storage_key: str) -> Path:
    root = Path(settings.document_storage_dir).resolve()
    path = (root / storage_key).resolve()
    if path.parent != root:
        raise ValueError("Invalid document storage key.")
    return path


def _store_file(data: bytes, suffix: str) -> str:
    root = Path(settings.document_storage_dir)
    try:
        root.mkdir(parents=True, exist_ok=True)
        descriptor, staged_name = tempfile.mkstemp(prefix=".upload-", dir=root)
        staged_path = Path(staged_name)
        try:
            with os.fdopen(descriptor, "wb") as staged_file:
                staged_file.write(data)
                staged_file.flush()
                os.fsync(staged_file.fileno())
            storage_key = f"{uuid.uuid4().hex}{suffix}"
            os.replace(staged_path, _storage_path(storage_key))
            return storage_key
        finally:
            staged_path.unlink(missing_ok=True)
    except OSError as exc:
        raise ApiError(
            503,
            "PROCESSING_UNAVAILABLE",
            "Document storage is temporarily unavailable.",
        ) from exc


def _remove_file(storage_key: str) -> None:
    _storage_path(storage_key).unlink(missing_ok=True)


def _fingerprint(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def create_document(
    db: Session,
    user_id: uuid.UUID,
    filename: str,
    media_type: str,
    source: BinaryIO,
    idempotency_key: str | None = None,
) -> tuple[Document, ProcessingJob]:
    safe_name = _safe_filename(filename)
    data = _read_bounded(source)
    normalized_media_type = _validate_upload(safe_name, media_type, data)
    content_hash = _fingerprint(data)

    if idempotency_key:
        existing_job = db.scalar(
            select(ProcessingJob).where(
                ProcessingJob.user_id == user_id,
                ProcessingJob.idempotency_key == idempotency_key,
            )
        )
        if existing_job is not None:
            if (
                existing_job.job_type != INGEST_JOB
                or existing_job.request_hash != content_hash
                or existing_job.document_id is None
            ):
                raise ApiError(
                    409,
                    "IDEMPOTENCY_CONFLICT",
                    "The idempotency key was already used for a different request.",
                )
            document = db.scalar(
                select(Document).where(
                    Document.id == existing_job.document_id,
                    Document.user_id == user_id,
                )
            )
            if document is not None:
                return document, existing_job

    storage_key = _store_file(
        data,
        Path(safe_name.replace("\\", "/")).suffix.lower(),
    )
    document = Document(
        id=uuid.uuid4(),
        user_id=user_id,
        filename=safe_name,
        storage_key=storage_key,
        media_type=normalized_media_type,
        status="PROCESSING",
    )
    job = ProcessingJob(
        user_id=user_id,
        document_id=document.id,
        job_type=INGEST_JOB,
        status="QUEUED",
        max_attempts=settings.document_job_max_attempts,
        idempotency_key=idempotency_key,
        request_hash=content_hash if idempotency_key else None,
    )
    db.add_all([document, job])
    record_learning_event(
        db,
        user_id,
        "DOCUMENT_UPLOADED",
        document_id=document.id,
        document_title=document.filename,
        details={"status": document.status},
    )
    try:
        db.commit()
        db.refresh(document)
        db.refresh(job)
    except IntegrityError as exc:
        db.rollback()
        _remove_file(storage_key)
        if idempotency_key:
            existing_job = db.scalar(
                select(ProcessingJob).where(
                    ProcessingJob.user_id == user_id,
                    ProcessingJob.idempotency_key == idempotency_key,
                )
            )
            if existing_job and existing_job.request_hash == content_hash:
                existing_document = db.get(Document, existing_job.document_id)
                if existing_document is not None:
                    return existing_document, existing_job
            raise ApiError(
                409,
                "IDEMPOTENCY_CONFLICT",
                "The idempotency key was already used for a different request.",
            ) from exc
        raise
    except Exception:
        db.rollback()
        _remove_file(storage_key)
        raise
    return document, job


def _text_chunks(pages: list[tuple[int | None, str]]) -> list[ExtractedChunk]:
    chunks: list[ExtractedChunk] = []
    for page_number, text in pages:
        normalized = " ".join(text.split())
        if not normalized:
            continue
        while normalized:
            if len(normalized) <= CHUNK_MAX_CHARACTERS:
                piece, normalized = normalized, ""
            else:
                boundary = normalized.rfind(" ", 0, CHUNK_MAX_CHARACTERS + 1)
                if boundary <= 0:
                    boundary = CHUNK_MAX_CHARACTERS
                piece, normalized = normalized[:boundary], normalized[boundary:].lstrip()
            chunks.append(
                ExtractedChunk(
                    chunk_index=len(chunks),
                    content=piece,
                    page_number=page_number,
                )
            )
    return chunks


def extract_document(
    content: bytes,
    filename: str,
    media_type: str,
) -> ExtractedDocument:
    suffix = Path(filename.replace("\\", "/")).suffix.lower()
    if MEDIA_TYPES.get(suffix) != media_type:
        raise DocumentProcessingError("Unsupported document type.")

    pages: list[tuple[int | None, str]] = []
    page_count: int | None = None
    try:
        if suffix == ".pdf":
            reader = PdfReader(BytesIO(content), strict=True)
            if len(reader.pages) > settings.document_max_pages:
                raise DocumentProcessingError("Document has too many pages.")
            page_count = len(reader.pages)
            pages = [
                (page_number, page.extract_text() or "")
                for page_number, page in enumerate(reader.pages, start=1)
            ]
        elif suffix == ".docx":
            from docx import Document as WordDocument

            word_document = WordDocument(BytesIO(content))
            parts = [paragraph.text for paragraph in word_document.paragraphs]
            parts.extend(
                cell.text
                for table in word_document.tables
                for row in table.rows
                for cell in row.cells
            )
            pages = [(None, "\n".join(parts))]
        else:
            from pptx import Presentation

            presentation = Presentation(BytesIO(content))
            if len(presentation.slides) > settings.document_max_pages:
                raise DocumentProcessingError("Document has too many slides.")
            page_count = len(presentation.slides)
            pages = [
                (
                    slide_number,
                    "\n".join(
                        shape.text
                        for shape in slide.shapes
                        if shape.has_text_frame
                    ),
                )
                for slide_number, slide in enumerate(
                    presentation.slides,
                    start=1,
                )
            ]
    except DocumentProcessingError:
        raise
    except (
        PdfReadError,
        OSError,
        ValueError,
        KeyError,
        zipfile.BadZipFile,
        DocxPackageNotFoundError,
        PptxPackageNotFoundError,
        XMLSyntaxError,
    ) as exc:
        raise DocumentProcessingError("The document could not be processed.") from exc
    except ImportError as exc:
        raise DocumentProcessingError("Document processing is unavailable.") from exc

    chunks = _text_chunks(pages)
    if not chunks:
        raise DocumentProcessingError("The document contains no extractable text.")
    return ExtractedDocument(page_count=page_count, chunks=chunks)


def _set_job_failure(
    db: Session,
    job: ProcessingJob,
    document: Document | None,
    failure_code: str,
) -> None:
    job.status = "FAILED"
    job.last_error_code = failure_code
    job.lease_expires_at = None
    if document is not None and document.status == "PROCESSING":
        document.status = "FAILED"
        document.failure_code = failure_code
    db.commit()


def _process_ingest(db: Session, job: ProcessingJob) -> None:
    document = db.get(Document, job.document_id)
    if document is None or document.status != "PROCESSING":
        job.status = "CANCELLED"
        job.lease_expires_at = None
        db.commit()
        return
    try:
        content = _storage_path(document.storage_key).read_bytes()
        extracted = extract_document(content, document.filename, document.media_type)
    except DocumentProcessingError:
        _set_job_failure(db, job, document, "DOCUMENT_UNREADABLE")
        return
    except AIProviderError as exc:
        _set_job_failure(db, job, document, exc.code)
        return
    except OSError:
        _set_job_failure(db, job, document, "DOCUMENT_STORAGE_ERROR")
        return

    db.execute(
        Quiz.__table__.delete().where(
            Quiz.document_id == document.id,
            Quiz.user_id == job.user_id,
        )
    )
    db.execute(
        DocumentChunk.__table__.delete().where(
            DocumentChunk.document_id == document.id,
            DocumentChunk.user_id == job.user_id,
        )
    )
    persisted_chunks = [
        DocumentChunk(
            user_id=job.user_id,
            document_id=document.id,
            chunk_index=chunk.chunk_index,
            content=chunk.content,
            page_number=chunk.page_number,
        )
        for chunk in extracted.chunks
    ]
    try:
        embed_document_chunks(db, persisted_chunks)
    except AIProviderError as exc:
        db.rollback()
        _set_job_failure(db, job, document, exc.code)
        return
    state_update = db.execute(
        update(Document)
        .where(
            Document.id == document.id,
            Document.user_id == job.user_id,
            Document.status == "PROCESSING",
        )
        .values(
            status="READY",
            failure_code=None,
            page_count=extracted.page_count,
            updated_at=datetime.now(timezone.utc),
        )
    )
    if state_update.rowcount != 1:
        job.status = "CANCELLED"
        job.lease_expires_at = None
        db.commit()
        return
    db.add_all(persisted_chunks)
    job.status = "SUCCEEDED"
    job.last_error_code = None
    job.lease_expires_at = None
    db.commit()


def _process_delete(db: Session, job: ProcessingJob) -> None:
    document = db.get(Document, job.document_id)
    if document is None:
        job.status = "SUCCEEDED"
        job.lease_expires_at = None
        db.commit()
        return
    if document.status not in {"DELETING", "DELETED"}:
        job.status = "CANCELLED"
        job.lease_expires_at = None
        db.commit()
        return
    try:
        _remove_file(document.storage_key)
    except OSError:
        if job.attempt_count >= job.max_attempts:
            job.status = "FAILED"
            job.last_error_code = "DOCUMENT_STORAGE_ERROR"
        else:
            job.status = "RETRY_WAIT"
            job.run_after = datetime.now(timezone.utc) + timedelta(
                seconds=2**job.attempt_count
            )
        job.lease_expires_at = None
        db.commit()
        return
    db.execute(
        DocumentChunk.__table__.delete().where(
            DocumentChunk.document_id == document.id,
            DocumentChunk.user_id == job.user_id,
        )
    )
    summaries = db.scalars(
        select(Summary).where(
            Summary.document_id == document.id,
            Summary.user_id == job.user_id,
        )
    ).all()
    for summary in summaries:
        summary.status = "DELETED"
        summary.content = None
    document.status = "DELETED"
    document.failure_code = None
    job.status = "SUCCEEDED"
    job.last_error_code = None
    job.lease_expires_at = None
    db.commit()


def process_next_job(db: Session) -> bool:
    now = datetime.now(timezone.utc)
    expired_jobs = db.scalars(
        select(ProcessingJob).where(
            ProcessingJob.status == "RUNNING",
            ProcessingJob.lease_expires_at <= now,
        )
    ).all()
    for expired_job in expired_jobs:
        expired_job.lease_expires_at = None
        if expired_job.attempt_count >= expired_job.max_attempts:
            expired_job.status = "FAILED"
            expired_job.last_error_code = "PROCESSING_ATTEMPTS_EXHAUSTED"
            document = db.get(Document, expired_job.document_id)
            if document is not None and document.status == "PROCESSING":
                document.status = "FAILED"
                document.failure_code = expired_job.last_error_code
            summary = db.get(Summary, expired_job.summary_id)
            if summary is not None and summary.status != "DELETED":
                summary.status = "FAILED"
                summary.failure_code = expired_job.last_error_code
            quiz = db.get(Quiz, expired_job.quiz_id)
            if quiz is not None and quiz.status != "DELETED":
                quiz.status = "FAILED"
                quiz.failure_code = expired_job.last_error_code
        else:
            expired_job.status = "RETRY_WAIT"
            expired_job.run_after = now + timedelta(seconds=2**expired_job.attempt_count)
            summary = db.get(Summary, expired_job.summary_id)
            if summary is not None and summary.status == "PROCESSING":
                summary.status = "QUEUED"
            quiz = db.get(Quiz, expired_job.quiz_id)
            if quiz is not None and quiz.status == "PROCESSING":
                quiz.status = "QUEUED"
    if expired_jobs:
        db.commit()

    job = db.scalar(
        select(ProcessingJob)
        .where(
            ProcessingJob.status.in_(("QUEUED", "RETRY_WAIT")),
            ProcessingJob.run_after <= now,
        )
        .order_by(ProcessingJob.created_at, ProcessingJob.id)
        .with_for_update(skip_locked=True)
    )
    if job is None:
        return False
    job.status = "RUNNING"
    job.attempt_count += 1
    job.lease_expires_at = now + timedelta(
        seconds=settings.document_job_lease_seconds
    )
    db.commit()
    if job.job_type == INGEST_JOB:
        _process_ingest(db, job)
    elif job.job_type == DELETE_JOB:
        _process_delete(db, job)
    elif job.job_type == "SUMMARY_GENERATE":
        from app.services.learning import process_summary_job

        process_summary_job(db, job)
    elif job.job_type == "QUIZ_GENERATE":
        from app.services.quizzes import process_quiz_job

        process_quiz_job(db, job)
    else:
        job.status = "FAILED"
        job.last_error_code = "UNKNOWN_JOB_TYPE"
        job.lease_expires_at = None
        db.commit()
    return True


def create_retry_job(
    db: Session,
    user_id: uuid.UUID,
    document_id: uuid.UUID,
    idempotency_key: str | None = None,
) -> ProcessingJob:
    request_hash = hashlib.sha256(f"retry:{document_id}".encode()).hexdigest()
    if idempotency_key:
        existing_job = db.scalar(
            select(ProcessingJob).where(
                ProcessingJob.user_id == user_id,
                ProcessingJob.idempotency_key == idempotency_key,
            )
        )
        if existing_job is not None:
            if (
                existing_job.job_type == INGEST_JOB
                and existing_job.document_id == document_id
                and existing_job.request_hash == request_hash
            ):
                return existing_job
            raise ApiError(
                409,
                "IDEMPOTENCY_CONFLICT",
                "The idempotency key was already used for a different request.",
            )

    document = db.scalar(
        select(Document).where(
            Document.id == document_id,
            Document.user_id == user_id,
        )
    )
    if document is None:
        raise ApiError(404, "NOT_FOUND", "Document not found.")
    if document.status != "FAILED":
        raise ApiError(
            409,
            "STATE_CONFLICT",
            "This document cannot be retried in its current state.",
        )
    document.status = "PROCESSING"
    document.failure_code = None
    job = ProcessingJob(
        user_id=user_id,
        document_id=document.id,
        job_type=INGEST_JOB,
        status="QUEUED",
        max_attempts=settings.document_job_max_attempts,
        idempotency_key=idempotency_key,
        request_hash=request_hash if idempotency_key else None,
    )
    db.add(job)
    try:
        db.commit()
        db.refresh(job)
    except IntegrityError as exc:
        db.rollback()
        raise ApiError(
            409,
            "IDEMPOTENCY_CONFLICT",
            "The idempotency key was already used for a different request.",
        ) from exc
    return job


def request_document_deletion(
    db: Session,
    user_id: uuid.UUID,
    document_id: uuid.UUID,
) -> tuple[Document, ProcessingJob | None]:
    document = db.scalar(
        select(Document).where(
            Document.id == document_id,
            Document.user_id == user_id,
        )
    )
    if document is None:
        raise ApiError(404, "NOT_FOUND", "Document not found.")
    if document.status == "DELETED":
        return document, None
    if document.status == "DELETING":
        job = db.scalar(
            select(ProcessingJob).where(
                ProcessingJob.user_id == user_id,
                ProcessingJob.document_id == document.id,
                ProcessingJob.job_type == DELETE_JOB,
                ProcessingJob.status.in_(("QUEUED", "RUNNING", "RETRY_WAIT")),
            )
        )
        if job is None:
            job = ProcessingJob(
                user_id=user_id,
                document_id=document.id,
                job_type=DELETE_JOB,
                status="QUEUED",
                max_attempts=settings.document_job_max_attempts,
            )
            db.add(job)
            db.commit()
            db.refresh(job)
        return document, job

    document.status = "DELETING"
    document.failure_code = None
    db.execute(
        update(ProcessingJob)
        .where(
            ProcessingJob.document_id == document.id,
            ProcessingJob.user_id == user_id,
            ProcessingJob.status.in_(("QUEUED", "RETRY_WAIT")),
        )
        .values(status="CANCELLED", updated_at=datetime.now(timezone.utc))
    )
    job = ProcessingJob(
        user_id=user_id,
        document_id=document.id,
        job_type=DELETE_JOB,
        status="QUEUED",
        max_attempts=settings.document_job_max_attempts,
    )
    db.add(job)
    db.commit()
    db.refresh(document)
    db.refresh(job)
    return document, job
